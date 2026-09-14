import streamlit as st
from openai import OpenAI
from docx import Document
import io

st.set_page_config(
    page_title="AI Assessment Designer",
    page_icon="🎓",
    layout="wide"
)

st.title("🎓 AI Assessment Designer")
st.subheader("AI-powered assessment creation assistant for university lecturers and professors")

api_key = st.sidebar.text_input(
    "Enter OpenAI API Key",
    type="password"
)

subject = st.text_input("Subject / Course Name")

level = st.selectbox(
    "Student Level",
    ["Undergraduate", "Postgraduate", "PhD"]
)

assessment_type = st.selectbox(
    "Assessment Type",
    [
        "MCQs",
        "Short Questions",
        "Long Questions",
        "Case Study",
        "Assignment",
        "Research Questions",
        "Mixed Assessment"
    ]
)

topics = st.text_area("Topics / Chapters")

bloom = st.multiselect(
    "Bloom's Taxonomy Level",
    [
        "Remember",
        "Understand",
        "Apply",
        "Analyze",
        "Evaluate",
        "Create"
    ],
    default=["Understand", "Apply", "Analyze"]
)

difficulty = st.selectbox(
    "Difficulty Level",
    ["Easy", "Moderate", "Advanced"]
)

number_questions = st.slider(
    "Number of Questions",
    1,
    50,
    10
)

marks = st.number_input(
    "Total Marks",
    value=100
)

learning_outcomes = st.text_area(
    "Course Learning Outcomes (Optional)"
)

def create_prompt():
    return f"""
You are an expert university assessment developer.

Create a high-quality assessment.

Course:
{subject}

Student Level:
{level}

Topics:
{topics}

Assessment Type:
{assessment_type}

Bloom Taxonomy Levels:
{bloom}

Difficulty:
{difficulty}

Number of Questions:
{number_questions}

Total Marks:
{marks}

Learning Outcomes:
{learning_outcomes}

Requirements:
1. Create academically valid questions.
2. Map each question with Bloom Taxonomy.
3. Provide marks distribution.
4. Provide answer key.
5. Provide marking rubric.
6. Align assessment with learning outcomes.

Format:
Question
Bloom Level
Marks
Expected Answer
Evaluation Criteria
"""

if st.button("🚀 Generate Assessment"):

    if not api_key:
        st.error("Please enter OpenAI API Key")
    else:
        client = OpenAI(api_key=api_key)

        with st.spinner("Creating assessment..."):

            response = client.chat.completions.create(
                model="gpt-4.1-mini",
                messages=[
                    {
                        "role": "system",
                        "content": "You are a senior university assessment specialist."
                    },
                    {
                        "role": "user",
                        "content": create_prompt()
                    }
                ]
            )

            st.session_state.assessment = response.choices[0].message.content


if "assessment" in st.session_state:

    st.divider()
    st.subheader("Generated Assessment")

    st.write(st.session_state.assessment)

    st.download_button(
        "Download TXT",
        st.session_state.assessment,
        file_name="assessment.txt"
    )

    if st.button("Create Word File"):

        doc = Document()
        doc.add_heading("AI Generated Assessment", level=1)
        doc.add_paragraph(st.session_state.assessment)

        buffer = io.BytesIO()
        doc.save(buffer)

        st.download_button(
            "Download Word",
            buffer.getvalue(),
            file_name="Assessment.docx"
        )
